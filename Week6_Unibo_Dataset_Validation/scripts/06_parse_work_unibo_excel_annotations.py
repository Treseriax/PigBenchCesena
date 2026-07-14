from pathlib import Path
import re
import json
import csv
from datetime import datetime, timedelta
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path.home() / "PigBench"
W6 = PROJECT_ROOT / "Week6_Unibo_Dataset_Validation"
OUT_GT = W6 / "outputs/unified_ground_truth"
OUT_STATS = W6 / "outputs/dataset_statistics"
NOTES = W6 / "notes"

OUT_GT.mkdir(parents=True, exist_ok=True)
OUT_STATS.mkdir(parents=True, exist_ok=True)
NOTES.mkdir(parents=True, exist_ok=True)

EXCEL_PATH = Path("/work/pig/datasets/Unibo/Giorno 1 - 22_7_2021 tlc1 FASCIA 9-10.xlsx")
VIDEO_NAMING_PATH = OUT_STATS / "unibo_raw_video_naming_pattern_analysis.csv"


def clean_text(x):
    if x is None:
        return ""
    s = str(x)
    s = s.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    s = " ".join(s.split())
    return s.strip()


def safe_to_csv(df, path):
    df.to_csv(
        path,
        index=False,
        quoting=csv.QUOTE_ALL,
        escapechar="\\",
        lineterminator="\n"
    )


def parse_date_from_filename(path: Path):
    # Giorno 1 - 22_7_2021 ...
    m = re.search(r"(\d{1,2})[_\-](\d{1,2})[_\-](20\d{2})", path.name)
    if m:
        day, month, year = map(int, m.groups())
        return datetime(year, month, day).date().isoformat()
    return ""


def normalize_colour(raw):
    r = clean_text(raw).lower()

    if "verde" in r or "green" in r:
        return "green"
    if "blu" in r or "blue" in r:
        return "blue"
    if "viola" in r or "purple" in r:
        return "purple"
    if "rosso testa" in r or "red neck" in r:
        return "red_neck"
    if "rosso coda" in r or "red tail" in r:
        return "red_tail"
    if "#" in r or "no color" in r or "no colour" in r:
        return "no_color"

    return r.replace(" ", "_")


def parse_time_range(s):
    s = clean_text(s)
    # 07,00-08,00
    m = re.search(r"(\d{1,2})[,.:](\d{2})\s*-\s*(\d{1,2})[,.:](\d{2})", s)
    if not m:
        return "", ""

    h1, m1, h2, m2 = m.groups()
    return f"{int(h1):02d}:{int(m1):02d}", f"{int(h2):02d}:{int(m2):02d}"


def build_behaviour_map(ws):
    # Definitions are stored in rows 25-37:
    # col A: Italian-English label, col D: code
    behaviour_map = {}

    for row in range(1, ws.max_row + 1):
        label = clean_text(ws.cell(row, 1).value)
        code = clean_text(ws.cell(row, 4).value)

        if code and label and code.isupper() and len(code) <= 5:
            behaviour_map[code] = label

    # Fallback if any code is missing.
    fallback = {
        "PI": "In piedi inattivi - Standing inactive",
        "SI": "Seduti inattivi - Sitting inactive",
        "LAI": "Stesi inattivi laterali - Laying in a lateral position",
        "STI": "Stesi inattivi sternali - Laying in a sternal position",
        "NU": "Si nutrono - Eating",
        "BE": "Bevono - Drinking",
        "DE": "Deambulano - Walking",
        "AN": "Annusano, grufolano - Sniffing - rooting",
        "IN": "INTERAZIONE NEUTRA - Neutral interaction",
        "IA": "INTERAZIONE AGGRESSIVA - Aggressive interaction",
        "MC": "MORDERE CODA - Tail biting",
        "ARR": "INTERAZIONE CON ARRICCHIMENTO - Interaction with the enrichment",
        "BOX": "INTERAZIONE BOX - Interaction with pen equipment",
    }

    for k, v in fallback.items():
        behaviour_map.setdefault(k, v)

    return behaviour_map


def load_tlc_video_lookup():
    if not VIDEO_NAMING_PATH.exists():
        return {}

    vids = pd.read_csv(VIDEO_NAMING_PATH)

    lookup = {}

    for _, r in vids.iterrows():
        family = str(r.get("naming_family", ""))
        if family != "TLC_B_time_range":
            continue

        cam = clean_text(r.get("parsed_camera_id", ""))
        pen = clean_text(r.get("parsed_pen_or_box_id", ""))
        start = clean_text(r.get("parsed_start_time", ""))
        end = clean_text(r.get("parsed_end_time", ""))

        key = (cam, pen, start, end)
        lookup[key] = {
            "video_path": r.get("absolute_path", ""),
            "filename": r.get("filename", ""),
            "fps": r.get("fps", ""),
            "frame_count": r.get("frame_count", ""),
            "duration_sec": r.get("duration_sec", ""),
        }

    return lookup


if not EXCEL_PATH.exists():
    raise FileNotFoundError(EXCEL_PATH)

wb = load_workbook(EXCEL_PATH, data_only=True)
ws = wb["TLC 1 big"]

date_iso = parse_date_from_filename(EXCEL_PATH)
camera_id = "TLC1"
pen_id = "B1"

behaviour_map = build_behaviour_map(ws)
video_lookup = load_tlc_video_lookup()

# Excel layout discovered from detailed inspection.
# block_name, hour_row, minute_row, data_rows
blocks = [
    {
        "block_name": "morning_07_13",
        "hour_row": 3,
        "minute_row": 4,
        "data_rows": list(range(6, 12)),
        "hour_start_columns": [2, 8, 14, 20, 26, 32],
    },
    {
        "block_name": "afternoon_13_19",
        "hour_row": 14,
        "minute_row": 15,
        "data_rows": list(range(17, 23)),
        "hour_start_columns": [2, 8, 14, 20, 26, 32],
    },
]

records = []

for block in blocks:
    hour_row = block["hour_row"]
    minute_row = block["minute_row"]
    hour_start_columns = block["hour_start_columns"]

    # Each hour block spans six columns: 0,10,20,30,40,50.
    col_to_hour = {}

    for hour_col in hour_start_columns:
        hour_label = clean_text(ws.cell(hour_row, hour_col).value)
        hour_start, hour_end = parse_time_range(hour_label)

        for offset in range(6):
            col = hour_col + offset
            col_to_hour[col] = {
                "hour_label": hour_label,
                "hour_start": hour_start,
                "hour_end": hour_end,
            }

    for row in block["data_rows"]:
        colour_raw = clean_text(ws.cell(row, 1).value)
        colour_id = normalize_colour(colour_raw)

        for col in range(2, 38):
            behaviour_code = clean_text(ws.cell(row, col).value)

            if not behaviour_code:
                continue

            hinfo = col_to_hour.get(col, {})
            hour_start = hinfo.get("hour_start", "")
            hour_end = hinfo.get("hour_end", "")
            hour_label = hinfo.get("hour_label", "")

            minute_value = clean_text(ws.cell(minute_row, col).value)

            try:
                minute_offset = int(float(minute_value))
            except Exception:
                minute_offset = None

            observation_time = ""
            nominal_next_observation_time = ""

            if date_iso and hour_start and minute_offset is not None:
                base_dt = datetime.fromisoformat(f"{date_iso}T{hour_start}:00")
                obs_dt = base_dt + timedelta(minutes=minute_offset)
                observation_time = obs_dt.isoformat(timespec="seconds")
                nominal_next_observation_time = (obs_dt + timedelta(minutes=10)).isoformat(timespec="seconds")

            # Link to TLC video if matching hourly file exists.
            key = (camera_id, pen_id, hour_start, hour_end)
            video_info = video_lookup.get(key, {})

            matched_video_path = video_info.get("video_path", "")
            matched_video_filename = video_info.get("filename", "")
            matched_video_status = "matched_tlc_hour_video" if matched_video_path else "no_matching_tlc_hour_video"

            fps = video_info.get("fps", "")
            frame_index_estimate = ""

            if matched_video_path and minute_offset is not None:
                try:
                    fps_float = float(fps)
                    frame_index_estimate = int(round(minute_offset * 60 * fps_float))
                except Exception:
                    frame_index_estimate = ""

            behaviour_label = behaviour_map.get(behaviour_code, "")

            records.append({
                "source_file": str(EXCEL_PATH),
                "sheet": ws.title,
                "block_name": block["block_name"],
                "date": date_iso,
                "camera_id": camera_id,
                "pen_id": pen_id,
                "colour_raw": colour_raw,
                "colour_id": colour_id,
                "pig_id": colour_id,
                "hour_label": hour_label,
                "hour_start": hour_start,
                "hour_end": hour_end,
                "observation_minute_in_hour": minute_offset,
                "observation_time": observation_time,
                "nominal_next_observation_time": nominal_next_observation_time,
                "behaviour_code": behaviour_code,
                "behaviour_label": behaviour_label,
                "label_source": "manual_excel_scan_sampling",
                "excel_row": row,
                "excel_column": col,
                "excel_cell": f"{get_column_letter(col)}{row}",
                "matched_video_path": matched_video_path,
                "matched_video_filename": matched_video_filename,
                "matched_video_status": matched_video_status,
                "matched_video_fps": fps,
                "frame_index_estimate": frame_index_estimate,
                "bbox_available": False,
                "bbox_source": "",
                "notes": "Excel scan-sampling label; bbox will be linked later from detector/tracker outputs where available.",
            })

ann = pd.DataFrame(records)

long_csv = OUT_GT / "work_unibo_excel_annotations_long.csv"
long_json = OUT_GT / "work_unibo_excel_annotations_long.json"
safe_to_csv(ann, long_csv)
long_json.write_text(json.dumps(records, indent=2, ensure_ascii=False))

behaviour_dist = (
    ann.groupby(["behaviour_code", "behaviour_label"])
    .size()
    .reset_index(name="count")
    .sort_values("count", ascending=False)
)

behaviour_dist["percentage"] = (behaviour_dist["count"] / len(ann) * 100).round(2)

behaviour_dist_path = OUT_GT / "work_unibo_excel_behaviour_distribution.csv"
safe_to_csv(behaviour_dist, behaviour_dist_path)

colour_dist = (
    ann.groupby(["colour_id", "colour_raw"])
    .size()
    .reset_index(name="count")
    .sort_values("colour_id")
)

colour_dist_path = OUT_GT / "work_unibo_excel_colour_distribution.csv"
safe_to_csv(colour_dist, colour_dist_path)

video_linkage = (
    ann.groupby(["hour_start", "hour_end", "matched_video_status", "matched_video_filename"])
    .size()
    .reset_index(name="annotation_count")
    .sort_values(["hour_start", "matched_video_status"])
)

video_linkage_path = OUT_GT / "work_unibo_excel_video_linkage_summary.csv"
safe_to_csv(video_linkage, video_linkage_path)

missing_rows = []

missing_rows.append({
    "check": "total_annotations",
    "value": len(ann),
    "issue": "",
})

missing_rows.append({
    "check": "missing_behaviour_label_mapping",
    "value": int((ann["behaviour_label"].astype(str) == "").sum()),
    "issue": "Behaviour code not found in extracted/fallback mapping",
})

missing_rows.append({
    "check": "annotations_without_matching_tlc_video",
    "value": int((ann["matched_video_status"] != "matched_tlc_hour_video").sum()),
    "issue": "No TLC hourly MP4 found for that hour block",
})

missing_rows.append({
    "check": "annotations_with_matching_tlc_video",
    "value": int((ann["matched_video_status"] == "matched_tlc_hour_video").sum()),
    "issue": "",
})

missing_rows.append({
    "check": "bbox_available_at_this_stage",
    "value": int(ann["bbox_available"].sum()),
    "issue": "This Excel parser extracts labels only; bbox linkage is a later step",
})

missing_df = pd.DataFrame(missing_rows)
missing_path = OUT_GT / "work_unibo_excel_missing_linkage_report.csv"
safe_to_csv(missing_df, missing_path)

note_path = NOTES / "work_unibo_excel_annotation_parser_notes.md"

with open(note_path, "w") as f:
    f.write("# Work Unibo Excel Annotation Parser Notes\n\n")

    f.write("## Purpose\n\n")
    f.write(
        "This parser converts the manual Excel scan-sampling annotation sheet into a long-format table. "
        "Each row corresponds to one pig colour/ID at one observation timestamp with one behaviour code.\n\n"
    )

    f.write("## Input file\n\n")
    f.write(f"- `{EXCEL_PATH}`\n\n")

    f.write("## Parsed structure\n\n")
    f.write(
        "- Rows 3-4 define the 07:00-13:00 hour blocks and 10-minute observation points.\n"
        "- Rows 6-11 contain pig-colour behaviour annotations for 07:00-13:00.\n"
        "- Rows 14-15 define the 13:00-19:00 hour blocks and 10-minute observation points.\n"
        "- Rows 17-22 contain pig-colour behaviour annotations for 13:00-19:00.\n\n"
    )

    f.write("## Behaviour distribution\n\n")
    f.write(behaviour_dist.to_markdown(index=False))
    f.write("\n\n")

    f.write("## Colour distribution\n\n")
    f.write(colour_dist.to_markdown(index=False))
    f.write("\n\n")

    f.write("## Video linkage summary\n\n")
    f.write(video_linkage.to_markdown(index=False))
    f.write("\n\n")

    f.write("## Missing / linkage report\n\n")
    f.write(missing_df.to_markdown(index=False))
    f.write("\n\n")

    f.write("## Interpretation\n\n")
    f.write(
        "The Excel file provides manual scan-sampling behaviour labels for six colour-coded pigs across 12 hours. "
        "Hourly TLC videos are currently matched for the available TLC1 B1 video files. "
        "Annotations without a matching TLC video are still preserved in the label table, but they cannot yet be directly visualized or linked to frame indices unless the corresponding raw videos are identified.\n"
    )

print("Saved:")
print(long_csv)
print(long_json)
print(behaviour_dist_path)
print(colour_dist_path)
print(video_linkage_path)
print(missing_path)
print(note_path)

print()
print("=== Parsed annotation count ===")
print(len(ann))

print()
print("=== Behaviour distribution ===")
print(behaviour_dist.to_string(index=False))

print()
print("=== Colour distribution ===")
print(colour_dist.to_string(index=False))

print()
print("=== Video linkage summary ===")
print(video_linkage.to_string(index=False))

print()
print("=== Missing/linkage report ===")
print(missing_df.to_string(index=False))

print()
print("=== First 20 parsed annotations ===")
cols = [
    "observation_time",
    "camera_id",
    "pen_id",
    "colour_id",
    "behaviour_code",
    "behaviour_label",
    "matched_video_status",
    "matched_video_filename",
    "frame_index_estimate",
    "excel_cell",
]
print(ann[cols].head(20).to_string(index=False))
