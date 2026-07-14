from pathlib import Path
import zipfile
import re
import html
import csv
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path.home() / "PigBench"
W6 = PROJECT_ROOT / "Week6_Unibo_Dataset_Validation"
OUT_GT = W6 / "outputs/unified_ground_truth"
NOTES = W6 / "notes"

OUT_GT.mkdir(parents=True, exist_ok=True)
NOTES.mkdir(parents=True, exist_ok=True)

EXCEL_PATH = Path("/work/pig/datasets/Unibo/Giorno 1 - 22_7_2021 tlc1 FASCIA 9-10.xlsx")
ETHOGRAM_PATH = Path("/work/pig/datasets/Unibo/spiegazione voci etogramma.docx")


def clean_text(x):
    if x is None:
        return ""
    s = str(x)
    s = s.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    s = " ".join(s.split())
    return s


def safe_to_csv(df, path):
    df.to_csv(
        path,
        index=False,
        quoting=csv.QUOTE_ALL,
        escapechar="\\",
        lineterminator="\n"
    )


if not EXCEL_PATH.exists():
    raise FileNotFoundError(EXCEL_PATH)

# 1) Inspect workbook structure with openpyxl.
wb = load_workbook(EXCEL_PATH, data_only=True)
sheet_rows = []

for ws in wb.worksheets:
    sheet_rows.append({
        "sheet_name": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "merged_ranges": "; ".join(str(rng) for rng in ws.merged_cells.ranges),
    })

sheet_df = pd.DataFrame(sheet_rows)
sheet_summary_path = OUT_GT / "work_unibo_excel_sheet_summary.csv"
safe_to_csv(sheet_df, sheet_summary_path)

# 2) Non-empty cells with row/column coordinates.
cell_rows = []

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            value_clean = clean_text(value)
            if value_clean == "":
                continue

            cell_rows.append({
                "sheet_name": ws.title,
                "row": cell.row,
                "column": cell.column,
                "coordinate": cell.coordinate,
                "value": value_clean,
            })

cells_df = pd.DataFrame(cell_rows)
nonempty_path = OUT_GT / "work_unibo_excel_nonempty_cells.csv"
safe_to_csv(cells_df, nonempty_path)

# 3) Raw rectangular table dump per sheet.
raw_dump_paths = []

for ws in wb.worksheets:
    matrix = []
    for r in range(1, ws.max_row + 1):
        row_values = []
        for c in range(1, ws.max_column + 1):
            row_values.append(clean_text(ws.cell(r, c).value))
        matrix.append(row_values)

    raw_df = pd.DataFrame(matrix)
    raw_path = OUT_GT / f"work_unibo_excel_raw_sheet_{ws.title.replace(' ', '_')}.csv"
    safe_to_csv(raw_df, raw_path)
    raw_dump_paths.append(raw_path)

# 4) Human-readable row summary: only rows with any content.
row_summary_rows = []

for ws in wb.worksheets:
    for r in range(1, ws.max_row + 1):
        values = []
        nonempty_count = 0

        for c in range(1, ws.max_column + 1):
            v = clean_text(ws.cell(r, c).value)
            if v:
                nonempty_count += 1
                values.append(f"C{c}={v}")

        if nonempty_count:
            row_summary_rows.append({
                "sheet_name": ws.title,
                "row": r,
                "nonempty_count": nonempty_count,
                "row_values": " | ".join(values),
            })

row_summary_df = pd.DataFrame(row_summary_rows)
row_summary_path = OUT_GT / "work_unibo_excel_nonempty_row_summary.csv"
safe_to_csv(row_summary_df, row_summary_path)

# 5) Extract DOCX text without relying on python-docx.
ethogram_text = ""
ethogram_status = ""

if ETHOGRAM_PATH.exists():
    try:
        with zipfile.ZipFile(ETHOGRAM_PATH) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
            # Convert Word XML paragraphs/runs roughly to text.
            xml = xml.replace("</w:p>", "\n")
            text = re.sub(r"<[^>]+>", "", xml)
            text = html.unescape(text)
            ethogram_text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
            ethogram_status = "extracted_docx_text"
    except Exception as e:
        ethogram_status = f"docx_extract_error: {type(e).__name__}: {e}"
else:
    ethogram_status = "ethogram_docx_not_found"

ethogram_txt_path = OUT_GT / "work_unibo_ethogram_extracted_text.txt"
ethogram_txt_path.write_text(ethogram_text if ethogram_text else ethogram_status)

# 6) Try to detect behaviour code explanations from ethogram text.
behaviour_codes = ["PI", "SI", "LAI", "STI", "NU", "BE", "DE", "AN", "IN", "IA", "MC", "ARR", "BOX", "SOMMA"]

ethogram_rows = []

for code in behaviour_codes:
    pattern = re.compile(rf"\b{re.escape(code)}\b(.{{0,250}})", flags=re.I)
    matches = pattern.findall(ethogram_text)
    ethogram_rows.append({
        "behaviour_code": code,
        "matches_found": len(matches),
        "matched_context_preview": " || ".join(clean_text(m) for m in matches[:3]),
    })

ethogram_df = pd.DataFrame(ethogram_rows)
ethogram_code_path = OUT_GT / "work_unibo_ethogram_code_contexts.csv"
safe_to_csv(ethogram_df, ethogram_code_path)

# 7) Build inspection note.
note_path = NOTES / "work_unibo_excel_and_ethogram_detailed_inspection.md"

with open(note_path, "w") as f:
    f.write("# Work Unibo Excel and Ethogram Detailed Inspection\n\n")

    f.write("## Purpose\n\n")
    f.write(
        "This note performs a detailed inspection of the Excel annotation/metadata file and the ethogram DOCX found directly in `/work/pig/datasets/Unibo`. "
        "The Excel file appears to use a manual observation layout with multiple hour blocks and behaviour codes rather than a simple flat table.\n\n"
    )

    f.write("## Excel file\n\n")
    f.write(f"- `{EXCEL_PATH}`\n\n")

    f.write("## Sheet summary\n\n")
    f.write(sheet_df.to_markdown(index=False))
    f.write("\n\n")

    f.write("## Non-empty row summary\n\n")
    if len(row_summary_df):
        f.write(row_summary_df.to_markdown(index=False))
    else:
        f.write("No non-empty rows found.")
    f.write("\n\n")

    f.write("## Behaviour code contexts from ethogram\n\n")
    f.write(ethogram_df.to_markdown(index=False))
    f.write("\n\n")

    f.write("## Ethogram extraction status\n\n")
    f.write(ethogram_status)
    f.write("\n\n")

    f.write("## Interpretation\n\n")
    f.write(
        "The next step is to design a parser for the Excel layout. The parser should convert the hour-block/minute-level manual annotations into a long-format table with timestamp, observation interval, pig/colour identifier if available, behaviour code, and behaviour label. "
        "The ethogram text should be used to map short behaviour codes to human-readable labels.\n"
    )

print("Saved:")
print(sheet_summary_path)
print(nonempty_path)
print(row_summary_path)
print(ethogram_txt_path)
print(ethogram_code_path)
print(note_path)

print()
print("=== Excel sheet summary ===")
print(sheet_df.to_string(index=False))

print()
print("=== Non-empty row summary ===")
print(row_summary_df.to_string(index=False))

print()
print("=== Ethogram code contexts ===")
print(ethogram_df.to_string(index=False))

print()
print("=== Ethogram text preview ===")
print(ethogram_text[:3000] if ethogram_text else ethogram_status)
