from pathlib import Path
from docx import Document

docx_path = Path("/work/pig/datasets/Unibo/spiegazione voci etogramma.docx")

print("Ethogram file:", docx_path)
print("Exists:", docx_path.exists())

doc = Document(docx_path)

print("\nParagraphs:")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"[P{i}] {text}")

print("\nTables:")
for ti, table in enumerate(doc.tables):
    print(f"\n--- Table {ti} ---")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
        print(" || ".join(cells))
