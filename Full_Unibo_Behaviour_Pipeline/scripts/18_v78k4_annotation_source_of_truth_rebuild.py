from pathlib import Path
from datetime import datetime
import re
import csv
import json
import hashlib
import zipfile
import pandas as pd
import openpyxl

ROOT = Path.home() / "PigBench"
FULL = ROOT / "Full_Unibo_Behaviour_Pipeline"
EXCEL_DIR = Path("/work/pig/datasets/Unibo/excel")

OUT = FULL / "outputs" / "v78k4_annotation_source_of_truth_rebuild"
PKG = OUT / "Full_Unibo_Annotation_Source_of_Truth_Rebuild"
NOTES = FULL / "notes"
REPORTS = FULL / "reports"
PROGRESS = FULL / "progress"

for p in [OUT, PKG, NOTES, REPORTS, PROGRESS]:
    p.mkdir(parents=True, exist_ok=True)

OUT_SHEETS = PKG / "v78k4_excel_sheet_source_of_truth_inventory.csv"
OUT_DAY1 = PKG / "v78k4_2021_07_22_annotation_target_map.csv"
OUT_RULES = PKG / "v78k4_annotation_interpretation_rules.csv"
OUT_DECODED_CANDIDATES = PKG / "v78k4_decoded_annotation_file_candidates.csv"
OUT_JOINED = PKG / "v78k4_decoded_annotations_joined_with_sheet_truth.csv"
OUT_SUMMARY = PKG / "v78k4_annotation_truth_summary.csv"
OUT_ISSUES = OUT / "v78k4_issues.csv"
OUT_DECISION = OUT / "v78k4_decision_summary.csv"
OUT_README = PKG / "README_v78k4_Annotation_Source_of_Truth_Rebuild.md"
OUT_MANIFEST = PKG / "v78k4_manifest.json"
OUT_ZIP = OUT / "Full_Unibo_Annotation_Source_of_Truth_Rebuild.zip"
OUT_SHA = OUT / "Full_Unibo_Annotation_Source_of_Truth_Rebuild.sha256"
OUT_NOTE = NOTES / "v78k4_annotation_source_of_truth_rebuild_notes.md"
OUT_REPORT = REPORTS / "v78k4_annotation_source_of_truth_rebuild_report.md"
OUT_PROGRESS = PROGRESS / "full_unibo_pipeline_progress_log.csv"

def clean(x):
    if x is None:
        return ""
    s = str(x).strip()
    if s.lower() in {"nan", "none", "null"}:
        return ""
    return s

def to_csv(df, path):
    df.to_csv(path, index=False, quoting=csv.QUOTE_ALL, escapechar="\\", lineterminator="\n")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def parse_date_from_filename(name):
    m = re.search(r"(\d{1,2})_(\d{1,2})_(\d{4})", name)
    if not m:
        return ""
    d, mo, y = m.groups()
    return f"{int(y):04d}-{int(mo):02d}-{int(d):02d}"

def parse_sheet_title(title):
    title = clean(title)
    low = title.lower()

    tlc_from_title = ""
    m = re.search(r"tlc\s*([1-6])", title, re.I)
    if m:
        tlc_from_title = f"TLC{m.group(1)}"

    sheet_alias = ""
    if "big" in low:
        sheet_alias = "big"
    elif "small" in low:
        sheet_alias = "small"

    explicit_room_pen = ""
    m2 = re.search(r"\b([BCM]\s*[0-9]+)\b", title, re.I)
    if m2:
        explicit_room_pen = m2.group(1).replace(" ", "").upper()

    return tlc_from_title, sheet_alias, explicit_room_pen

def parse_row1_metadata(ws):
    vals = [clean(c.value) for c in ws[1]]
    vals = [v for v in vals if v]
    row1_text = " ; ".join(vals)

    camera_num = ""
    room = ""
    pen_num = ""

    for i, v in enumerate(vals):
        low = v.lower()
        if low == "camera" and i + 1 < len(vals):
            camera_num = clean(vals[i + 1])
        elif low == "room" and i + 1 < len(vals):
            room = clean(vals[i + 1]).upper()
        elif low == "pen" and i + 1 < len(vals):
            pen_num = clean(vals[i + 1])

    tlc = f"TLC{camera_num}" if camera_num else ""
    room_pen = f"{room}{pen_num}" if room and pen_num else ""

    return row1_text, camera_num, tlc, room, pen_num, room_pen

def detect_time_headers(ws, max_rows=20, max_cols=80):
    hits = []
    time_re = re.compile(r"\b([01]?\d|2[0-3])[:.][0-5]\d\b|\b([01]?\d|2[0-3])00\b")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), max_col=min(ws.max_column, max_cols)):
        for cell in row:
            v = clean(cell.value)
            if v and time_re.search(v):
                hits.append(f"{cell.coordinate}={v}")
    return ";".join(hits[:50])

def sample_nonempty_grid(ws, max_rows=12, max_cols=20):
    lines = []
    for r in range(1, min(ws.max_row, max_rows) + 1):
        vals = []
        for c in range(1, min(ws.max_column, max_cols) + 1):
            v = clean(ws.cell(r, c).value)
            vals.append(v)
        if any(vals):
            lines.append(" | ".join(vals))
    return "\n".join(lines)

issues = []
sheet_rows = []

excel_files = sorted(list(EXCEL_DIR.glob("*.xlsx")) + list(EXCEL_DIR.glob("*.xlsm")))

for path in excel_files:
    date = parse_date_from_filename(path.name)

    try:
        wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
        for ws in wb.worksheets:
            tlc_from_title, sheet_alias, explicit_room_pen = parse_sheet_title(ws.title)
            row1_text, camera_num, row_tlc, room, pen_num, row_room_pen = parse_row1_metadata(ws)

            if not row_tlc and not tlc_from_title and not row_room_pen:
                continue

            source_tlc = row_tlc or tlc_from_title
            source_room_pen = row_room_pen or explicit_room_pen

            if sheet_alias and row_room_pen:
                annotation_unit_type = "sheet_alias_resolved_by_row1"
            elif explicit_room_pen and row_room_pen:
                annotation_unit_type = "explicit_sheet_pen_confirmed_by_row1"
            elif row_room_pen:
                annotation_unit_type = "row1_camera_room_pen_metadata"
            else:
                annotation_unit_type = "unresolved_sheet_metadata"

            sheet_rows.append({
                "date": date,
                "excel_file_name": path.name,
                "excel_file": str(path),
                "sheet_name": ws.title,
                "sheet_state": ws.sheet_state,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "row1_text": row1_text,
                "camera_num": camera_num,
                "tlc_camera": source_tlc,
                "room": room,
                "pen_num": pen_num,
                "resolved_room_pen": source_room_pen,
                "sheet_alias": sheet_alias,
                "sheet_explicit_room_pen": explicit_room_pen,
                "annotation_unit_type": annotation_unit_type,
                "source_truth_statement": f"{ws.title} is annotation source for {source_tlc} / {source_room_pen}",
                "time_header_hits_sample": detect_time_headers(ws),
                "grid_sample_first_rows": sample_nonempty_grid(ws),
            })
        wb.close()
    except Exception as e:
        issues.append({
            "item": str(path),
            "issue_type": "excel_read_error",
            "issue_detail": str(e),
            "severity": "warning",
        })

sheets = pd.DataFrame(sheet_rows)
to_csv(sheets, OUT_SHEETS)

if sheets.empty:
    issues.append({
        "item": "excel_sheet_inventory",
        "issue_type": "hard_no_sheet_truth_rows",
        "issue_detail": "No Excel sheet metadata rows were extracted.",
        "severity": "hard",
    })

# 2021-07-22 target map: source-of-truth sheet interpretation.
day1 = sheets[sheets["date"] == "2021-07-22"].copy()
day1 = day1[[
    "date", "tlc_camera", "resolved_room_pen", "sheet_name",
    "sheet_alias", "sheet_explicit_room_pen", "row1_text",
    "annotation_unit_type", "source_truth_statement"
]].sort_values(["tlc_camera", "resolved_room_pen", "sheet_name"])
to_csv(day1, OUT_DAY1)

rules = pd.DataFrame([
    {
        "rule_id": "R1",
        "rule": "The Excel sheet is the annotation unit.",
        "interpretation": "If a sheet resolves to TLC2 / B4, all behaviour rows decoded from that sheet belong to TLC2 / B4 unless another explicit row-level override exists.",
    },
    {
        "rule_id": "R2",
        "rule": "big/small are aliases, not pen names.",
        "interpretation": "Resolve big/small using row-1 metadata Camera / Room / Pen. Example: TLC1 big on 2021-07-22 resolves to B1.",
    },
    {
        "rule_id": "R3",
        "rule": "Do not infer frame position from Excel sheet labels.",
        "interpretation": "Excel tells which pen the annotation describes. It does not say where that pen is inside the image frame.",
    },
    {
        "rule_id": "R4",
        "rule": "Camera/pen association and video/ROI mapping are separate.",
        "interpretation": "Annotation source truth can be complete even when encoded c-code video mapping or frame-level ROI remains unresolved.",
    },
])
to_csv(rules, OUT_RULES)

# Find decoded annotation candidate files from prior stages.
candidate_paths = []
for p in FULL.rglob("*.csv"):
    name = p.name.lower()
    full = str(p).lower()
    if ("v77c" in full or "layout_aware" in full or "decoded" in name) and ("annotation" in name or "window" in name):
        candidate_paths.append(p)

cand = pd.DataFrame([{
    "path": str(p),
    "name": p.name,
    "size_bytes": p.stat().st_size,
} for p in sorted(candidate_paths)])
to_csv(cand, OUT_DECODED_CANDIDATES)

# Try to join best decoded annotation file with sheet source truth, if columns allow.
joined = pd.DataFrame()
join_status = "no_decoded_join_attempted"

if candidate_paths:
    best = sorted(candidate_paths, key=lambda x: x.stat().st_size, reverse=True)[0]
    try:
        dec = pd.read_csv(best).fillna("")
        for c in dec.columns:
            if dec[c].dtype == object:
                dec[c] = dec[c].map(clean)

        sheet_col = None
        for c in dec.columns:
            if c.lower() in {"sheet_name", "sheet", "worksheet", "excel_sheet"}:
                sheet_col = c
                break

        date_col = None
        for c in dec.columns:
            if c.lower() in {"date", "annotation_date", "day"}:
                date_col = c
                break

        if sheet_col:
            meta_cols = [
                "date", "sheet_name", "tlc_camera", "resolved_room_pen",
                "room", "pen_num", "sheet_alias", "annotation_unit_type",
                "source_truth_statement"
            ]
            meta = sheets[meta_cols].drop_duplicates()

            if date_col:
                joined = dec.merge(
                    meta,
                    left_on=[date_col, sheet_col],
                    right_on=["date", "sheet_name"],
                    how="left",
                    suffixes=("", "_sheet_truth")
                )
                join_status = f"joined_on_{date_col}_and_{sheet_col}"
            else:
                joined = dec.merge(
                    meta.drop(columns=["date"]).drop_duplicates(),
                    left_on=sheet_col,
                    right_on="sheet_name",
                    how="left",
                    suffixes=("", "_sheet_truth")
                )
                join_status = f"joined_on_{sheet_col}_only"

            to_csv(joined, OUT_JOINED)
        else:
            join_status = f"decoded_file_found_but_no_sheet_column: {best}"
    except Exception as e:
        join_status = f"decoded_join_error: {e}"

if joined.empty:
    # still create an empty file with explanation
    pd.DataFrame([{
        "join_status": join_status,
        "note": "No decoded annotation table could be joined automatically. Sheet source-of-truth inventory is still valid."
    }]).to_csv(OUT_JOINED, index=False, quoting=csv.QUOTE_ALL)

# Summaries.
summary_rows = []
if not sheets.empty:
    for (date, tlc), g in sheets.groupby(["date", "tlc_camera"]):
        pens = ";".join(sorted(set([clean(x) for x in g["resolved_room_pen"].tolist() if clean(x)])))
        aliases = []
        for _, r in g.iterrows():
            if clean(r["sheet_alias"]):
                aliases.append(f"{r['sheet_alias']}={r['resolved_room_pen']}")
        summary_rows.append({
            "date": date,
            "tlc_camera": tlc,
            "resolved_room_pens": pens,
            "aliases": ";".join(sorted(set(aliases))),
            "source_sheet_count": len(g),
            "source_sheets": ";".join(sorted(set(g["sheet_name"].tolist()))),
        })
summary = pd.DataFrame(summary_rows).sort_values(["date", "tlc_camera"]) if summary_rows else pd.DataFrame()
to_csv(summary, OUT_SUMMARY)

# QA for Day 1.
expected_day1 = {
    "TLC1": {"B1", "B3"},
    "TLC2": {"B4", "B6"},
    "TLC3": {"C1", "C3"},
    "TLC4": {"C4", "C6"},
    "TLC5": {"M1", "M2"},
    "TLC6": {"M3", "M4"},
}

if not day1.empty:
    for tlc, expected in expected_day1.items():
        obs = set(day1[day1["tlc_camera"] == tlc]["resolved_room_pen"].tolist())
        missing = expected - obs
        if missing:
            issues.append({
                "item": f"2021-07-22 {tlc}",
                "issue_type": "hard_day1_expected_annotation_targets_missing",
                "issue_detail": f"Missing expected annotation target(s): {sorted(missing)}; observed={sorted(obs)}",
                "severity": "hard",
            })

issues.append({
    "item": "annotation_boundary",
    "issue_type": "info_annotation_source_truth_not_video_roi",
    "issue_detail": "Excel sheet metadata resolves annotation source truth. It does not define encoded c-code video mapping or frame-level spatial ROI.",
    "severity": "info",
})

issues_df = pd.DataFrame(issues, columns=["item", "issue_type", "issue_detail", "severity"])
to_csv(issues_df, OUT_ISSUES)

hard_count = int((issues_df["severity"] == "hard").sum()) if len(issues_df) else 0
warning_count = int((issues_df["severity"] == "warning").sum()) if len(issues_df) else 0
info_count = int((issues_df["severity"] == "info").sum()) if len(issues_df) else 0

day1_ok = hard_count == 0 and not day1.empty
decision_label = "annotation_source_truth_rebuilt" if day1_ok else "annotation_source_truth_has_blocking_issues"

readme = f"""# v78k4 Annotation Source-of-Truth Rebuild

## Purpose

This stage treats Excel as the annotation source of truth.

It answers:

- Which sheet is the annotation unit?
- Which TLC / Camera / Room / Pen does each sheet describe?
- How should `big` and `small` be interpreted?
- Which parts are annotation truth and which parts are video/spatial mapping?

## Key rule

Do not look for B1/B3 visually in the frame during annotation decoding.

The sheet itself defines the annotation target.

Example:

- `TLC 1 big` resolves to `TLC1 / B1` using row-1 metadata.
- `TLC 1 B3` resolves to `TLC1 / B3`.
- `TLC 2 B4` resolves to `TLC2 / B4`.

## Boundary

This stage does not solve encoded c-code mapping and does not solve frame-level ROI.
"""

OUT_README.write_text(readme, encoding="utf-8")
OUT_REPORT.write_text(readme, encoding="utf-8")

manifest = {
    "version": "v78k4_annotation_source_of_truth_rebuild",
    "generated_at": datetime.now().isoformat(timespec="seconds"),
    "excel_files": len(excel_files),
    "sheet_truth_rows": int(len(sheets)),
    "day1_target_rows": int(len(day1)),
    "summary_rows": int(len(summary)),
    "decoded_candidate_files": int(len(cand)),
    "decoded_join_rows": int(len(joined)),
    "decoded_join_status": join_status,
    "hard_issue_count": hard_count,
    "warning_count": warning_count,
    "claim_boundary": "Excel annotation source truth only; no video mapping or ROI",
}
OUT_MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

if OUT_ZIP.exists():
    OUT_ZIP.unlink()

with zipfile.ZipFile(OUT_ZIP, "w", compression=zipfile.ZIP_DEFLATED) as z:
    for p in sorted(PKG.rglob("*")):
        if p.is_file():
            z.write(p, p.relative_to(OUT))

zip_hash = sha256_file(OUT_ZIP)
OUT_SHA.write_text(f"{zip_hash}  {OUT_ZIP.name}\n")

decision = pd.DataFrame([{
    "v78k4_decision": decision_label,
    "excel_files": len(excel_files),
    "sheet_truth_rows": int(len(sheets)),
    "day1_target_rows": int(len(day1)),
    "summary_rows": int(len(summary)),
    "decoded_candidate_files": int(len(cand)),
    "decoded_join_status": join_status,
    "decoded_join_rows": int(len(joined)),
    "zip_path": str(OUT_ZIP),
    "zip_sha256": zip_hash,
    "hard_issue_count": hard_count,
    "warning_count": warning_count,
    "info_count": info_count,
    "ready_for_annotation_truth_use": bool(hard_count == 0 and len(sheets) > 0),
    "ready_for_video_mapping": False,
    "ready_for_spatial_roi": False,
    "claim_scope": "annotation_source_truth_only",
    "generated_at": datetime.now().isoformat(timespec="seconds"),
}])
to_csv(decision, OUT_DECISION)

OUT_NOTE.write_text(
    "# v78k4 Annotation Source-of-Truth Rebuild\n\n"
    f"- Decision: {decision_label}\n"
    f"- Excel files: {len(excel_files)}\n"
    f"- Sheet truth rows: {len(sheets)}\n"
    f"- 2021-07-22 target rows: {len(day1)}\n"
    f"- Decoded candidate files: {len(cand)}\n"
    f"- Decoded join status: {join_status}\n"
    f"- Hard issues: {hard_count}\n"
    f"- Ready for annotation truth use: {bool(hard_count == 0 and len(sheets) > 0)}\n\n"
    "Core correction: B1/B3 are sheet-level annotation targets. Do not infer their frame position during annotation decoding.\n",
    encoding="utf-8"
)

progress_row = pd.DataFrame([{
    "date": datetime.now().date().isoformat(),
    "stage": "v78k4",
    "task_name": "Annotation source-of-truth rebuild",
    "status": "PASS" if hard_count == 0 else "NEEDS_FIX",
    "input_summary": "Raw Unibo Excel annotation files",
    "output_summary": str(PKG),
    "hard_issues": hard_count,
    "warnings": warning_count,
    "next_action": "Use sheet-level annotation truth before any video c-code or ROI mapping.",
}])

if OUT_PROGRESS.exists():
    old = pd.read_csv(OUT_PROGRESS)
    progress = pd.concat([old, progress_row], ignore_index=True)
else:
    progress = progress_row
to_csv(progress, OUT_PROGRESS)

print("=== v78k4 decision ===")
print(decision.to_string(index=False))

print("\n=== 2021-07-22 annotation target map ===")
print(day1.to_string(index=False) if not day1.empty else "none")

print("\n=== annotation interpretation rules ===")
print(rules.to_string(index=False))

print("\n=== summary ===")
print(summary.to_string(index=False) if not summary.empty else "none")

print("\n=== decoded candidates ===")
print(cand.to_string(index=False) if not cand.empty else "none")

print("\n=== issues ===")
print(issues_df.to_string(index=False))
