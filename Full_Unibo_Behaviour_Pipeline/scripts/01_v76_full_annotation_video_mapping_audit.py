from pathlib import Path
from datetime import datetime
import re
import csv
import json
import hashlib
import zipfile
import traceback
import pandas as pd


ROOT = Path.home() / "PigBench"
UNIBO = Path("/work/pig/datasets/Unibo")

FULL = ROOT / "Full_Unibo_Behaviour_Pipeline"
SCRIPTS = FULL / "scripts"
OUT = FULL / "outputs" / "v76_full_annotation_video_mapping_audit"
PKG = OUT / "Full_Unibo_Annotation_Video_Mapping_Audit"
NOTES = FULL / "notes"
REPORTS = FULL / "reports"
PROGRESS = FULL / "progress"

for p in [FULL, SCRIPTS, OUT, PKG, NOTES, REPORTS, PROGRESS]:
    p.mkdir(parents=True, exist_ok=True)

OUT_EXCEL_FILES = PKG / "v76_all_excel_files_inventory.csv"
OUT_SHEETS = PKG / "v76_all_excel_sheets_inventory.csv"
OUT_SHEET_PREVIEW = PKG / "v76_excel_sheet_content_preview.csv"
OUT_VIDEOS = PKG / "v76_all_videos_inventory.csv"
OUT_VIDEO_PARSE = PKG / "v76_video_name_parse_inventory.csv"
OUT_CANDIDATES = PKG / "v76_candidate_excel_video_matches.csv"
OUT_TOP_MATCHES = PKG / "v76_top_excel_video_matches.csv"
OUT_COVERAGE = PKG / "v76_coverage_summary.csv"
OUT_UNMATCHED_VIDEOS = PKG / "v76_unmatched_videos.csv"
OUT_UNMATCHED_SHEETS = PKG / "v76_unmatched_annotation_sheets.csv"
OUT_QA = PKG / "v76_quality_checks.csv"
OUT_README = PKG / "README_v76_Full_Unibo_Annotation_Video_Mapping_Audit.md"
OUT_MANIFEST = PKG / "v76_manifest.json"

OUT_DECISION = OUT / "v76_decision_summary.csv"
OUT_ISSUES = OUT / "v76_issues.csv"
OUT_ZIP = OUT / "Full_Unibo_Annotation_Video_Mapping_Audit.zip"
OUT_SHA = OUT / "Full_Unibo_Annotation_Video_Mapping_Audit.sha256"
OUT_NOTE = NOTES / "v76_full_annotation_video_mapping_audit_notes.md"
OUT_REPORT = REPORTS / "v76_full_annotation_video_mapping_audit_report.md"
OUT_PROGRESS = PROGRESS / "full_unibo_pipeline_progress_log.csv"


KNOWN_BEHAVIOURS = {"STI", "LAI", "BOX", "AN", "NU", "PI", "IN", "SI", "DE", "BE", "IA"}
KNOWN_IDENTITY_HINTS = {
    "blue", "green", "purple", "red", "red_neck", "red_tail",
    "cyan", "no_colour", "no_color", "pink", "colour", "color"
}


def safe_to_csv(df, path):
    df.to_csv(path, index=False, quoting=csv.QUOTE_ALL, escapechar="\\", lineterminator="\n")


def clean(x):
    if pd.isna(x):
        return ""
    s = str(x).strip()
    if s.lower() in ["nan", "none", "null"]:
        return ""
    return s


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def short_error(e):
    return str(e).replace("\n", " ").replace("\r", " ")[:1000]


def parse_day_date_from_excel_name(name):
    text = name

    day_num = ""
    m = re.search(r"Giorno\s*([0-9]+)", text, flags=re.IGNORECASE)
    if m:
        day_num = m.group(1)

    date_iso = ""
    m = re.search(r"([0-9]{1,2})_([0-9]{1,2})_([0-9]{4})", text)
    if m:
        d = int(m.group(1))
        mo = int(m.group(2))
        y = int(m.group(3))
        date_iso = f"{y:04d}-{mo:02d}-{d:02d}"

    return day_num, date_iso


def hhmm_to_minutes(hhmm):
    s = str(hhmm).strip()
    s = re.sub(r"[^0-9]", "", s)

    if len(s) <= 2:
        h = int(s)
        m = 0
    elif len(s) == 3:
        h = int(s[0])
        m = int(s[1:])
    else:
        h = int(s[:2])
        m = int(s[2:4])

    return h * 60 + m


def minutes_to_hhmm(m):
    if m == "" or m is None:
        return ""
    m = int(m)
    h = m // 60
    mm = m % 60
    return f"{h:02d}:{mm:02d}"


def parse_time_range_from_text(text):
    t = str(text)

    # Strong pattern: 0700-0800, 1000-1100
    m = re.search(r"(?<![0-9])([0-2]?[0-9][0-5][0-9])\s*[-–]\s*([0-2]?[0-9][0-5][0-9])(?![0-9])", t)
    if m:
        return hhmm_to_minutes(m.group(1)), hhmm_to_minutes(m.group(2)), "hhmm-hhmm"

    # FASCIA 9-10 or fascia 9 - 10
    m = re.search(r"FASCIA\s*([0-2]?[0-9])\s*[-–]\s*([0-2]?[0-9])", t, flags=re.IGNORECASE)
    if m:
        return int(m.group(1)) * 60, int(m.group(2)) * 60, "fascia-hour-hour"

    # TLC1 B1 800-900
    m = re.search(r"B\s*[0-9]\s+([0-2]?[0-9][0-5][0-9])\s*[-–]\s*([0-2]?[0-9][0-5][0-9])", t, flags=re.IGNORECASE)
    if m:
        return hhmm_to_minutes(m.group(1)), hhmm_to_minutes(m.group(2)), "b-hhmm-hhmm"

    return "", "", ""


def parse_camera_pen_from_text(text):
    t = str(text)

    camera = ""
    pen = ""

    m = re.search(r"TLC\s*([0-9]+)", t, flags=re.IGNORECASE)
    if m:
        camera = f"TLC{m.group(1)}"

    m = re.search(r"\bB\s*([0-9]+)\b", t, flags=re.IGNORECASE)
    if m:
        pen = f"B{m.group(1)}"

    return camera, pen


def parse_video_name(path):
    name = Path(path).stem
    full_name = Path(path).name

    row = {
        "video_path": str(path),
        "video_filename": full_name,
        "video_id": name,
        "video_name_type": "",
        "parsed_camera": "",
        "parsed_pen": "",
        "parsed_camera_code": "",
        "parsed_date": "",
        "parsed_start_minute": "",
        "parsed_end_minute_estimate": "",
        "parsed_start_hhmm": "",
        "parsed_end_hhmm_estimate": "",
        "parse_confidence": "low",
        "parse_notes": "",
    }

    # Human-friendly videos, e.g. TLC1 B1 1000-1100.mp4
    cam, pen = parse_camera_pen_from_text(name)
    ts, te, ttype = parse_time_range_from_text(name)

    if cam or pen or ttype:
        row["video_name_type"] = "friendly_tlc_time_range"
        row["parsed_camera"] = cam
        row["parsed_pen"] = pen
        row["parsed_start_minute"] = ts
        row["parsed_end_minute_estimate"] = te
        row["parsed_start_hhmm"] = minutes_to_hhmm(ts) if ts != "" else ""
        row["parsed_end_hhmm_estimate"] = minutes_to_hhmm(te) if te != "" else ""
        row["parse_confidence"] = "high" if cam and pen and ttype else "medium"
        row["parse_notes"] = "Parsed TLC/B/time range from filename."
        return row

    # Encoded videos, e.g. c0001210722150000 = c0001 + 210722 + 150000
    m = re.match(r"^c([0-9]{4})([0-9]{6})([0-9]{6})$", name, flags=re.IGNORECASE)
    if m:
        camera_code = "c" + m.group(1)
        yymmdd = m.group(2)
        hhmmss = m.group(3)

        yy = int(yymmdd[0:2])
        mo = int(yymmdd[2:4])
        dd = int(yymmdd[4:6])
        yyyy = 2000 + yy

        hh = int(hhmmss[0:2])
        mm = int(hhmmss[2:4])

        start = hh * 60 + mm

        row["video_name_type"] = "encoded_camera_datetime"
        row["parsed_camera_code"] = camera_code
        row["parsed_date"] = f"{yyyy:04d}-{mo:02d}-{dd:02d}"
        row["parsed_start_minute"] = start
        row["parsed_start_hhmm"] = minutes_to_hhmm(start)
        row["parse_confidence"] = "high"
        row["parse_notes"] = "Parsed camera code, date and start time from encoded filename."
        return row

    row["video_name_type"] = "unknown"
    row["parse_notes"] = "Could not parse video name."
    return row


def get_video_duration(path):
    try:
        import cv2
        cap = cv2.VideoCapture(str(path))
        if not cap.isOpened():
            return "", "", "cv2_open_failed"

        fps = cap.get(cv2.CAP_PROP_FPS)
        frames = cap.get(cv2.CAP_PROP_FRAME_COUNT)
        width = cap.get(cv2.CAP_PROP_FRAME_WIDTH)
        height = cap.get(cv2.CAP_PROP_FRAME_HEIGHT)
        cap.release()

        duration = ""
        if fps and fps > 0 and frames and frames > 0:
            duration = float(frames) / float(fps)

        return duration, f"{int(width)}x{int(height)}" if width and height else "", "ok"

    except Exception as e:
        return "", "", "duration_error:" + short_error(e)


def inspect_excel_workbook(path):
    try:
        import openpyxl
    except Exception as e:
        raise RuntimeError("openpyxl is required to inspect xlsx files: " + short_error(e))

    workbook_rows = []
    preview_rows = []

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    day_num, file_date = parse_day_date_from_excel_name(path.name)
    file_time_start, file_time_end, file_time_type = parse_time_range_from_text(path.name)
    file_cam, file_pen = parse_camera_pen_from_text(path.name)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        combined_hint = f"{path.name} {sheet_name}"
        sheet_time_start, sheet_time_end, sheet_time_type = parse_time_range_from_text(combined_hint)
        sheet_cam, sheet_pen = parse_camera_pen_from_text(combined_hint)

        if sheet_time_start == "" and file_time_start != "":
            sheet_time_start = file_time_start
            sheet_time_end = file_time_end
            sheet_time_type = file_time_type

        if not sheet_cam and file_cam:
            sheet_cam = file_cam
        if not sheet_pen and file_pen:
            sheet_pen = file_pen

        nonempty_cells = 0
        first_nonempty_row = ""
        last_nonempty_row = ""
        max_nonempty_col = 0
        behaviour_counts = {k: 0 for k in sorted(KNOWN_BEHAVIOURS)}
        identity_hint_counts = {k: 0 for k in sorted(KNOWN_IDENTITY_HINTS)}
        sample_values = []
        first_rows_text = []

        # Full scan can still be fine for these few xlsx files, but cap extremely large sheets defensively.
        row_limit = 5000

        for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if ridx > row_limit:
                break

            row_values = []
            row_has_value = False

            for cidx, value in enumerate(row, start=1):
                if value is None:
                    continue

                s = str(value).strip()
                if not s:
                    continue

                row_has_value = True
                nonempty_cells += 1

                if first_nonempty_row == "":
                    first_nonempty_row = ridx
                last_nonempty_row = ridx
                max_nonempty_col = max(max_nonempty_col, cidx)

                upper = s.upper().strip()
                if upper in behaviour_counts:
                    behaviour_counts[upper] += 1

                low = s.lower().strip()
                for hint in identity_hint_counts:
                    if hint in low:
                        identity_hint_counts[hint] += 1

                if len(sample_values) < 80:
                    sample_values.append(s[:120])

                if ridx <= 20:
                    row_values.append(s[:60])

            if ridx <= 20 and row_values:
                first_rows_text.append(" | ".join(row_values[:30]))

        sample_text = " ; ".join(sample_values[:60])
        first_rows_joined = "\n".join(first_rows_text[:20])

        workbook_rows.append({
            "excel_path": str(path),
            "excel_filename": path.name,
            "excel_sha256": sha256_file(path),
            "day_number_from_filename": day_num,
            "date_from_filename": file_date,
            "sheet_name": sheet_name,
            "sheet_max_row_reported": ws.max_row,
            "sheet_max_column_reported": ws.max_column,
            "scanned_row_limit": row_limit,
            "nonempty_cells_scanned": nonempty_cells,
            "first_nonempty_row": first_nonempty_row,
            "last_nonempty_row_scanned": last_nonempty_row,
            "max_nonempty_col_scanned": max_nonempty_col,
            "parsed_camera": sheet_cam,
            "parsed_pen": sheet_pen,
            "parsed_time_start_minute": sheet_time_start,
            "parsed_time_end_minute": sheet_time_end,
            "parsed_time_start_hhmm": minutes_to_hhmm(sheet_time_start) if sheet_time_start != "" else "",
            "parsed_time_end_hhmm": minutes_to_hhmm(sheet_time_end) if sheet_time_end != "" else "",
            "time_parse_type": sheet_time_type,
            "known_behaviour_total": sum(behaviour_counts.values()),
            "behaviour_counts_json": json.dumps(behaviour_counts, ensure_ascii=False),
            "identity_hint_counts_json": json.dumps(identity_hint_counts, ensure_ascii=False),
            "sample_text": sample_text,
        })

        preview_rows.append({
            "excel_path": str(path),
            "excel_filename": path.name,
            "sheet_name": sheet_name,
            "first_nonempty_rows_preview": first_rows_joined,
        })

    wb.close()
    return workbook_rows, preview_rows


def time_overlap(a_start, a_end, b_start, b_end):
    if a_start == "" or a_end == "" or b_start == "" or b_end == "":
        return False

    a_start = int(a_start)
    a_end = int(a_end)
    b_start = int(b_start)
    b_end = int(b_end)

    return max(a_start, b_start) < min(a_end, b_end)


def score_sheet_video(sheet, video):
    score = 0
    reasons = []

    sheet_date = clean(sheet.get("date_from_filename", ""))
    video_date = clean(video.get("parsed_date", ""))

    if sheet_date and video_date and sheet_date == video_date:
        score += 50
        reasons.append("date_match")

    # Friendly TLC names may not have parsed date; day/date match cannot be scored there.
    sheet_cam = clean(sheet.get("parsed_camera", ""))
    video_cam = clean(video.get("parsed_camera", ""))

    if sheet_cam and video_cam and sheet_cam == video_cam:
        score += 20
        reasons.append("camera_match")

    sheet_pen = clean(sheet.get("parsed_pen", ""))
    video_pen = clean(video.get("parsed_pen", ""))

    if sheet_pen and video_pen and sheet_pen == video_pen:
        score += 20
        reasons.append("pen_match")

    st = sheet.get("parsed_time_start_minute", "")
    se = sheet.get("parsed_time_end_minute", "")
    vt = video.get("parsed_start_minute", "")
    ve = video.get("parsed_end_minute_estimate", "")

    if st != "" and vt != "":
        if int(st) == int(vt):
            score += 25
            reasons.append("start_time_exact_match")

    if time_overlap(st, se, vt, ve):
        score += 30
        reasons.append("time_overlap")

    sheet_text = f"{sheet.get('excel_filename','')} {sheet.get('sheet_name','')}".lower()
    video_text = f"{video.get('video_filename','')}".lower()

    # Light token match for friendly videos.
    for token in ["tlc1", "b1", "0700", "0800", "900", "1000", "1100", "1200", "1300", "1400", "1500"]:
        if token in sheet_text and token in video_text:
            score += 4
            reasons.append(f"token_{token}")

    if score >= 90:
        quality = "high"
    elif score >= 60:
        quality = "medium"
    elif score >= 35:
        quality = "low"
    else:
        quality = "weak"

    return score, quality, ";".join(reasons)


issues = []

if not UNIBO.exists():
    issues.append({
        "item": str(UNIBO),
        "issue_type": "hard_unibo_dataset_missing",
        "issue_detail": "Unibo dataset root does not exist.",
        "severity": "hard",
    })

excel_files = []
video_files = []

if UNIBO.exists():
    excel_files = sorted(list(UNIBO.rglob("*.xlsx")) + list(UNIBO.rglob("*.xls")))
    video_files = sorted(
        list(UNIBO.rglob("*.mp4"))
        + list(UNIBO.rglob("*.avi"))
        + list(UNIBO.rglob("*.mov"))
        + list(UNIBO.rglob("*.mkv"))
    )

excel_file_rows = []

for p in excel_files:
    day_num, date_iso = parse_day_date_from_excel_name(p.name)
    f_time_start, f_time_end, f_time_type = parse_time_range_from_text(p.name)
    cam, pen = parse_camera_pen_from_text(p.name)

    excel_file_rows.append({
        "excel_path": str(p),
        "excel_filename": p.name,
        "size_mb": round(p.stat().st_size / (1024 * 1024), 3),
        "sha256": sha256_file(p),
        "day_number_from_filename": day_num,
        "date_from_filename": date_iso,
        "parsed_camera": cam,
        "parsed_pen": pen,
        "parsed_time_start_minute": f_time_start,
        "parsed_time_end_minute": f_time_end,
        "parsed_time_start_hhmm": minutes_to_hhmm(f_time_start) if f_time_start != "" else "",
        "parsed_time_end_hhmm": minutes_to_hhmm(f_time_end) if f_time_end != "" else "",
        "time_parse_type": f_time_type,
    })

excel_files_df = pd.DataFrame(excel_file_rows)
safe_to_csv(excel_files_df, OUT_EXCEL_FILES)

sheet_rows = []
preview_rows = []

for p in excel_files:
    try:
        rows, previews = inspect_excel_workbook(p)
        sheet_rows.extend(rows)
        preview_rows.extend(previews)
    except Exception as e:
        issues.append({
            "item": str(p),
            "issue_type": "warning_excel_inspection_failed",
            "issue_detail": short_error(e),
            "severity": "warning",
        })

sheets_df = pd.DataFrame(sheet_rows)
preview_df = pd.DataFrame(preview_rows)
safe_to_csv(sheets_df, OUT_SHEETS)
safe_to_csv(preview_df, OUT_SHEET_PREVIEW)

video_rows = []
parse_rows = []

for p in video_files:
    parsed = parse_video_name(p)
    duration_sec, resolution, duration_status = get_video_duration(p)

    end_est = parsed.get("parsed_end_minute_estimate", "")
    if end_est == "" and duration_sec != "" and parsed.get("parsed_start_minute", "") != "":
        end_est = int(parsed["parsed_start_minute"]) + int(round(float(duration_sec) / 60.0))
        parsed["parsed_end_minute_estimate"] = end_est
        parsed["parsed_end_hhmm_estimate"] = minutes_to_hhmm(end_est)

    row = {
        **parsed,
        "size_mb": round(p.stat().st_size / (1024 * 1024), 3),
        "duration_sec": duration_sec,
        "resolution": resolution,
        "duration_status": duration_status,
    }

    video_rows.append(row)
    parse_rows.append({
        "video_path": parsed["video_path"],
        "video_filename": parsed["video_filename"],
        "video_name_type": parsed["video_name_type"],
        "parsed_camera": parsed["parsed_camera"],
        "parsed_pen": parsed["parsed_pen"],
        "parsed_camera_code": parsed["parsed_camera_code"],
        "parsed_date": parsed["parsed_date"],
        "parsed_start_hhmm": parsed["parsed_start_hhmm"],
        "parsed_end_hhmm_estimate": parsed["parsed_end_hhmm_estimate"],
        "parse_confidence": parsed["parse_confidence"],
        "parse_notes": parsed["parse_notes"],
    })

videos_df = pd.DataFrame(video_rows)
parse_df = pd.DataFrame(parse_rows)
safe_to_csv(videos_df, OUT_VIDEOS)
safe_to_csv(parse_df, OUT_VIDEO_PARSE)

candidate_rows = []

for _, sheet in sheets_df.iterrows():
    for _, video in videos_df.iterrows():
        score, quality, reasons = score_sheet_video(sheet, video)

        if score <= 0:
            continue

        candidate_rows.append({
            "excel_path": sheet["excel_path"],
            "excel_filename": sheet["excel_filename"],
            "sheet_name": sheet["sheet_name"],
            "sheet_date": sheet["date_from_filename"],
            "sheet_camera": sheet["parsed_camera"],
            "sheet_pen": sheet["parsed_pen"],
            "sheet_time_start_hhmm": sheet["parsed_time_start_hhmm"],
            "sheet_time_end_hhmm": sheet["parsed_time_end_hhmm"],
            "video_path": video["video_path"],
            "video_filename": video["video_filename"],
            "video_name_type": video["video_name_type"],
            "video_date": video["parsed_date"],
            "video_camera": video["parsed_camera"],
            "video_pen": video["parsed_pen"],
            "video_camera_code": video["parsed_camera_code"],
            "video_start_hhmm": video["parsed_start_hhmm"],
            "video_end_hhmm_estimate": video["parsed_end_hhmm_estimate"],
            "match_score": score,
            "match_quality": quality,
            "match_reasons": reasons,
        })

candidates_df = pd.DataFrame(candidate_rows)

if len(candidates_df):
    candidates_df = candidates_df.sort_values(
        ["excel_filename", "sheet_name", "match_score"],
        ascending=[True, True, False],
    )

safe_to_csv(candidates_df, OUT_CANDIDATES)

if len(candidates_df):
    top_matches = (
        candidates_df.sort_values(["excel_filename", "sheet_name", "match_score"], ascending=[True, True, False])
        .groupby(["excel_path", "sheet_name"], as_index=False)
        .head(1)
        .reset_index(drop=True)
    )
else:
    top_matches = pd.DataFrame(columns=[
        "excel_path", "excel_filename", "sheet_name", "video_path", "video_filename",
        "match_score", "match_quality", "match_reasons"
    ])

safe_to_csv(top_matches, OUT_TOP_MATCHES)

matched_video_paths = set(candidates_df[candidates_df["match_quality"].isin(["high", "medium"])]["video_path"].tolist()) if len(candidates_df) else set()
matched_sheet_keys = set(
    zip(
        candidates_df[candidates_df["match_quality"].isin(["high", "medium"])]["excel_path"],
        candidates_df[candidates_df["match_quality"].isin(["high", "medium"])]["sheet_name"],
    )
) if len(candidates_df) else set()

unmatched_videos = videos_df[~videos_df["video_path"].isin(matched_video_paths)].copy() if len(videos_df) else pd.DataFrame()
safe_to_csv(unmatched_videos, OUT_UNMATCHED_VIDEOS)

if len(sheets_df):
    unmatched_sheet_mask = []
    for _, r in sheets_df.iterrows():
        unmatched_sheet_mask.append((r["excel_path"], r["sheet_name"]) not in matched_sheet_keys)
    unmatched_sheets = sheets_df[unmatched_sheet_mask].copy()
else:
    unmatched_sheets = pd.DataFrame()

safe_to_csv(unmatched_sheets, OUT_UNMATCHED_SHEETS)

duplicate_excel_sha_count = 0
if len(excel_files_df):
    duplicate_excel_sha_count = int(excel_files_df.duplicated("sha256", keep=False).sum())

coverage = pd.DataFrame([{
    "excel_file_count": int(len(excel_files_df)),
    "excel_duplicate_sha_rows": duplicate_excel_sha_count,
    "excel_sheet_count": int(len(sheets_df)),
    "video_count": int(len(videos_df)),
    "friendly_tlc_video_count": int((videos_df["video_name_type"] == "friendly_tlc_time_range").sum()) if len(videos_df) else 0,
    "encoded_camera_video_count": int((videos_df["video_name_type"] == "encoded_camera_datetime").sum()) if len(videos_df) else 0,
    "unknown_video_name_count": int((videos_df["video_name_type"] == "unknown").sum()) if len(videos_df) else 0,
    "candidate_match_count": int(len(candidates_df)),
    "high_candidate_match_count": int((candidates_df["match_quality"] == "high").sum()) if len(candidates_df) else 0,
    "medium_candidate_match_count": int((candidates_df["match_quality"] == "medium").sum()) if len(candidates_df) else 0,
    "low_candidate_match_count": int((candidates_df["match_quality"] == "low").sum()) if len(candidates_df) else 0,
    "matched_video_count_medium_or_high": int(len(matched_video_paths)),
    "unmatched_video_count_medium_or_high": int(len(unmatched_videos)),
    "matched_sheet_count_medium_or_high": int(len(matched_sheet_keys)),
    "unmatched_sheet_count_medium_or_high": int(len(unmatched_sheets)),
    "known_behaviour_total_in_scanned_sheets": int(sheets_df["known_behaviour_total"].sum()) if len(sheets_df) else 0,
}])

safe_to_csv(coverage, OUT_COVERAGE)

qa_rows = []

def add_qa(name, expected, actual, passed, severity, detail):
    qa_rows.append({
        "check_name": name,
        "expected": str(expected),
        "actual": str(actual),
        "passed": bool(passed),
        "severity": severity,
        "detail": detail,
    })

add_qa("unibo_root_exists", True, UNIBO.exists(), UNIBO.exists(), "hard", "Dataset root should exist.")
add_qa("excel_files_found", ">=4", len(excel_files_df), len(excel_files_df) >= 4, "hard", "Expected multiple Excel annotation files.")
add_qa("video_files_found", ">=80", len(videos_df), len(videos_df) >= 80, "hard", "Expected the full Unibo raw video set.")
add_qa("excel_sheets_found", ">0", len(sheets_df), len(sheets_df) > 0, "hard", "Excel workbook sheets should be inspectable.")
add_qa("known_behaviour_tokens_found", ">0", int(sheets_df["known_behaviour_total"].sum()) if len(sheets_df) else 0, (int(sheets_df["known_behaviour_total"].sum()) if len(sheets_df) else 0) > 0, "warning", "Known behaviour tokens help confirm annotation content.")
add_qa("candidate_matches_found", ">0", len(candidates_df), len(candidates_df) > 0, "hard", "Candidate Excel-video matches should be generated.")
add_qa("friendly_tlc_videos_found", ">0", int((videos_df["video_name_type"] == "friendly_tlc_time_range").sum()) if len(videos_df) else 0, (int((videos_df["video_name_type"] == "friendly_tlc_time_range").sum()) if len(videos_df) else 0) > 0, "warning", "Friendly TLC videos should be detected.")
add_qa("encoded_camera_videos_found", ">0", int((videos_df["video_name_type"] == "encoded_camera_datetime").sum()) if len(videos_df) else 0, (int((videos_df["video_name_type"] == "encoded_camera_datetime").sum()) if len(videos_df) else 0) > 0, "warning", "Encoded camera videos should be detected.")
add_qa("coverage_summary_created", 1, len(coverage), len(coverage) == 1, "hard", "Coverage summary should be created.")

qa = pd.DataFrame(qa_rows)
safe_to_csv(qa, OUT_QA)

hard_quality_failures = int(((qa["severity"] == "hard") & (~qa["passed"])).sum())
warning_quality_failures = int(((qa["severity"] == "warning") & (~qa["passed"])).sum())

if hard_quality_failures:
    issues.append({
        "item": "v76_quality_checks",
        "issue_type": "hard_full_mapping_audit_failed",
        "issue_detail": f"{hard_quality_failures} hard QA checks failed.",
        "severity": "hard",
    })

if warning_quality_failures:
    issues.append({
        "item": "v76_quality_checks",
        "issue_type": "warning_full_mapping_audit_has_warnings",
        "issue_detail": f"{warning_quality_failures} warning QA checks failed.",
        "severity": "warning",
    })

if len(unmatched_videos) > 0:
    issues.append({
        "item": "unmatched_videos",
        "issue_type": "info_unmatched_or_low_confidence_videos_exist",
        "issue_detail": f"{len(unmatched_videos)} videos currently have no medium/high candidate sheet match. This is expected in the first full audit and requires mapping review.",
        "severity": "info",
    })

if len(unmatched_sheets) > 0:
    issues.append({
        "item": "unmatched_sheets",
        "issue_type": "info_unmatched_or_low_confidence_sheets_exist",
        "issue_detail": f"{len(unmatched_sheets)} sheets currently have no medium/high candidate video match. This is expected in the first full audit and requires mapping review.",
        "severity": "info",
    })

issues.append({
    "item": "scope",
    "issue_type": "info_audit_only",
    "issue_detail": "v76 is a full annotation/video mapping audit. It does not train models, run tracking, or create final GT.",
    "severity": "info",
})

issues_df = pd.DataFrame(issues, columns=["item", "issue_type", "issue_detail", "severity"])
safe_to_csv(issues_df, OUT_ISSUES)

hard_issue_count = int((issues_df["severity"] == "hard").sum()) if len(issues_df) else 0
warning_count = int((issues_df["severity"] == "warning").sum()) if len(issues_df) else 0
info_count = int((issues_df["severity"] == "info").sum()) if len(issues_df) else 0

manifest = {
    "version": "v76_full_annotation_video_mapping_audit",
    "created_at": datetime.now().isoformat(timespec="seconds"),
    "unibo_root": str(UNIBO),
    "excel_file_count": int(len(excel_files_df)),
    "excel_sheet_count": int(len(sheets_df)),
    "video_count": int(len(videos_df)),
    "candidate_match_count": int(len(candidates_df)),
    "matched_video_count_medium_or_high": int(len(matched_video_paths)),
    "matched_sheet_count_medium_or_high": int(len(matched_sheet_keys)),
    "hard_issue_count": hard_issue_count,
    "claim_boundary": "inventory and candidate mapping audit only",
}

OUT_MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False))

readme_text = f"""# v76 Full Unibo Annotation + Video Mapping Audit

## Purpose

This is the first full-pipeline step after Week8.

It inventories all available Unibo Excel annotation files and all raw videos, inspects workbook sheets, parses video filenames, and creates candidate Excel-to-video mappings.

## Important boundary

This stage does not:
- train a model,
- run tracking,
- create final moving GT,
- assign final behaviour labels.

It only prepares the full annotation/video mapping foundation.

## Counts

- Excel files: {len(excel_files_df)}
- Excel sheets: {len(sheets_df)}
- Videos: {len(videos_df)}
- Candidate matches: {len(candidates_df)}
- Medium/high matched videos: {len(matched_video_paths)}
- Medium/high matched sheets: {len(matched_sheet_keys)}
- Hard issues: {hard_issue_count}

## Next stage

v77 should review and finalize the annotation-video mapping table, including manual resolution for ambiguous or unmatched cases.
"""

OUT_README.write_text(readme_text)
OUT_REPORT.write_text(readme_text)

if OUT_ZIP.exists():
    OUT_ZIP.unlink()

with zipfile.ZipFile(OUT_ZIP, "w", compression=zipfile.ZIP_DEFLATED) as z:
    for p in sorted(PKG.rglob("*")):
        if p.is_file():
            z.write(p, p.relative_to(OUT))

zip_hash = sha256_file(OUT_ZIP)
OUT_SHA.write_text(f"{zip_hash}  {OUT_ZIP.name}\n")

decision = pd.DataFrame([{
    "v76_decision": "full_annotation_video_mapping_audit_completed" if hard_issue_count == 0 else "full_annotation_video_mapping_audit_has_blocking_issues",
    "unibo_root": str(UNIBO),
    "excel_file_count": int(len(excel_files_df)),
    "excel_sheet_count": int(len(sheets_df)),
    "video_count": int(len(videos_df)),
    "candidate_match_count": int(len(candidates_df)),
    "high_candidate_match_count": int((candidates_df["match_quality"] == "high").sum()) if len(candidates_df) else 0,
    "medium_candidate_match_count": int((candidates_df["match_quality"] == "medium").sum()) if len(candidates_df) else 0,
    "matched_video_count_medium_or_high": int(len(matched_video_paths)),
    "unmatched_video_count_medium_or_high": int(len(unmatched_videos)),
    "matched_sheet_count_medium_or_high": int(len(matched_sheet_keys)),
    "unmatched_sheet_count_medium_or_high": int(len(unmatched_sheets)),
    "known_behaviour_total_in_scanned_sheets": int(sheets_df["known_behaviour_total"].sum()) if len(sheets_df) else 0,
    "zip_path": str(OUT_ZIP),
    "zip_sha256": zip_hash,
    "hard_quality_failures": hard_quality_failures,
    "warning_quality_failures": warning_quality_failures,
    "hard_issue_count": hard_issue_count,
    "warning_count": warning_count,
    "info_count": info_count,
    "issue_count": int(len(issues_df)),
    "ready_for_v77_mapping_resolution": bool(hard_issue_count == 0),
    "claim_scope": "full_annotation_video_inventory_and_candidate_mapping_audit_only",
    "generated_at": datetime.now().isoformat(timespec="seconds"),
}])

safe_to_csv(decision, OUT_DECISION)

OUT_NOTE.write_text(
    "# v76 Full Unibo Annotation + Video Mapping Audit\n\n"
    f"- v76 decision: {decision.iloc[0]['v76_decision']}\n"
    f"- Excel files: {len(excel_files_df)}\n"
    f"- Excel sheets: {len(sheets_df)}\n"
    f"- Videos: {len(videos_df)}\n"
    f"- Candidate matches: {len(candidates_df)}\n"
    f"- High matches: {decision.iloc[0]['high_candidate_match_count']}\n"
    f"- Medium matches: {decision.iloc[0]['medium_candidate_match_count']}\n"
    f"- Medium/high matched videos: {len(matched_video_paths)}\n"
    f"- Medium/high unmatched videos: {len(unmatched_videos)}\n"
    f"- Medium/high matched sheets: {len(matched_sheet_keys)}\n"
    f"- Medium/high unmatched sheets: {len(unmatched_sheets)}\n"
    f"- Known behaviour tokens in scanned sheets: {decision.iloc[0]['known_behaviour_total_in_scanned_sheets']}\n"
    f"- Hard issues: {hard_issue_count}\n"
    f"- Ready for v77 mapping resolution: {bool(hard_issue_count == 0)}\n\n"
    "This is a full inventory and candidate mapping audit. It does not create final GT, tracking, or model outputs.\n"
)

progress_row = pd.DataFrame([{
    "date": datetime.now().date().isoformat(),
    "stage": "v76",
    "task_name": "Full Unibo annotation and video mapping audit",
    "status": "PASS_WITH_WARNINGS" if hard_issue_count == 0 and warning_count > 0 else ("PASS" if hard_issue_count == 0 else "NEEDS_FIX"),
    "input_summary": str(UNIBO),
    "output_summary": str(PKG),
    "hard_issues": hard_issue_count,
    "warnings": warning_count,
    "next_action": "Resolve annotation-video mapping in v77." if hard_issue_count == 0 else "Fix v76 blocking issues.",
}])

if OUT_PROGRESS.exists():
    old = pd.read_csv(OUT_PROGRESS)
    progress = pd.concat([old, progress_row], ignore_index=True)
else:
    progress = progress_row

safe_to_csv(progress, OUT_PROGRESS)

print("=== v76 decision ===")
print(decision.to_string(index=False))

print("\n=== coverage summary ===")
print(coverage.to_string(index=False))

print("\n=== top candidate matches sample ===")
if len(top_matches):
    print(top_matches.head(30).to_string(index=False))
else:
    print("No top matches.")

print("\n=== QA ===")
print(qa.to_string(index=False))

print("\n=== issues ===")
print(issues_df.to_string(index=False) if len(issues_df) else "No issues found.")
